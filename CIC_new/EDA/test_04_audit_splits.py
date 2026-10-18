"""Focused tests for split isolation, episodes, boundaries and raw reconciliation."""
import csv
import gzip
import importlib.util
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest

import numpy as np

spec = importlib.util.spec_from_file_location("split_eda", Path(__file__).with_name("04_audit_splits.py"))
eda = importlib.util.module_from_spec(spec)
spec.loader.exec_module(eda)


def verify_saved_report(folder):
    """Read-only independent check of exported cached formulas against JSON."""
    import openpyxl
    folder=Path(folder)
    d=json.loads((folder/"results.json").read_text(encoding="utf-8"))
    wb=openpyxl.load_workbook(folder/"04_cic_split_audit.xlsx",data_only=True)
    try:
        for i,p in enumerate(d["candidates"],5):
            name=p["candidate"]
            row=[wb["Comparison"].cell(i,c).value for c in range(1,13)]
            assert row[0]==name
            counts=[r for r in d["composition"] if r["candidate"]==name]
            for j,r in enumerate(counts,1):
                assert abs(row[j]-r["sequences"]/d["metadata"]["source_sequences"])<1e-12
            types=[r for r in d["coverage"] if r["candidate"]==name]
            expected=[sum(r[s+"_windows"]>0 for r in types) for s in eda.SPLITS]
            expected += [sum(r[k] for r in types) for k in ("validation_unseen_in_train","test_unseen_in_train","present_in_all_three")]
            expected += [next(r["cross_episodes_300s"] for r in d["isolation"] if r["candidate"]==name)]
            assert row[5:]==expected,(name,row[5:],expected)
        for i,r in enumerate(d["composition"],5):
            observed=wb["Composition"].cell(i,5).value
            if r["sequences"]:
                assert abs(observed-r["positive_sequences"]/r["sequences"])<1e-12
            else:
                assert observed=="n.a."
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    assert cell.data_type!="e",(sheet.title,cell.coordinate,cell.value)
        assert len(wb.sheetnames)==8
    finally:
        wb.close()


class SplitTests(unittest.TestCase):
    @unittest.skipUnless(os.environ.get("CIC_SPLIT_EDA_REPORT"),"Set CIC_SPLIT_EDA_REPORT to verify a saved full report")
    def test_saved_workbook(self):
        verify_saved_report(os.environ["CIC_SPLIT_EDA_REPORT"])

    def test_episode_gap_and_midnight(self):
        bins = {("A",0): [1,9,2,{0}], ("A",6): [60_009,60_010,1,{1}],
                ("A",13): [130_011,130_012,1,{2}],
                ("A",8639): [86_399_999,86_399_999,1,{3}],
                ("A",8640): [86_400_000,86_400_000,1,{4}]}
        eps = eda.make_episodes(bins,60)
        self.assertEqual(len(eps),4)
        self.assertEqual(eps[0]["members"],{0,1})  # equality at threshold remains together
        self.assertEqual(sum(e["attack_flows"] for e in eps),6)
        self.assertEqual(len(eda.make_episodes(bins,900)),3)

    def test_purge_actual_ends_and_half_open_boundaries(self):
        v={"start":np.array([0,100_000,200_000,300_000,400_000]),
           "first":np.array([1,100_001,200_001,300_001,400_001]),
           "max_end":np.array([199_999,200_000,220_000,320_000,420_000])}
        original=np.array([0,0,1,1,2],dtype=np.uint8)
        kept,drops=eda.purge(v,original,0)
        self.assertEqual(kept.tolist(),[0,3,1,1,2])
        self.assertEqual(drops["interval_removed"],1)  # inclusive end exactly at boundary
        self.assertEqual(int(eda.overlaps_other_split(v,kept).sum()),0)
        guarded,_=eda.purge(v,original,100_000)
        self.assertEqual(guarded.tolist(),[3,3,3,3,3])
        mixed={k:np.array([x[0],x[0]]) for k,x in v.items()}
        with self.assertRaises(ValueError):
            eda.purge(mixed,np.array([0,1]),0)

    def test_guarded_blocks_union_transitively(self):
        t=np.array([0,1_800_000,3_600_000,5_400_000,7_200_000])
        eps=[{"first_ms":1,"last_ms":1_800_001}, {"first_ms":1_800_002,"last_ms":3_600_001}]
        groups=eda.guarded_blocks(t,eps,1_800_000)
        self.assertEqual(groups[0],groups[2])
        self.assertNotEqual(groups[2],groups[3])
        codes=eda.partition_groups(np.arange(5),np.array([60,10,10,10,10]))
        self.assertEqual(codes.tolist(),[0,1,1,2,2])

    def test_analysis_reconciles_and_saves_manifests(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            db=root/"raw.sqlite"
            conn=sqlite3.connect(db)
            conn.execute("CREATE TABLE network_flows(ipv4_src_addr TEXT,flow_start_milliseconds INTEGER,flow_end_milliseconds INTEGER,label INTEGER,attack TEXT)")
            catalog=root/"windows_100s.csv.gz"
            fields=["source_ip","start_ms","end_ms_exclusive","flows","attack_flows","first_flow_offset_s","last_flow_offset_s","attack_types"]
            rows=[]
            for source in ("A","B","C"):
                for day in range(10):
                    for block in range(4):
                        start=day*eda.DAY+block*3_600_000
                        positive=source=="A" and day%2==0
                        attack="Alpha" if positive else "Benign"
                        rows.append([source,start,start+eda.WINDOW,2,int(positive),.001,.005,"Alpha" if positive else ""])
                        conn.executemany("INSERT INTO network_flows VALUES(?,?,?,?,?)",
                            [(source,start+1,start+1_000,int(positive),attack), (source,start+5,start+5_000,0,"Benign")])
            conn.commit()
            conn.close()
            with gzip.open(catalog,"wt",encoding="utf-8-sig",newline="") as f:
                w=csv.writer(f)
                w.writerow(fields)
                w.writerows(rows)
            summary={"length_s":100,"windows":120,"flows":240,"attack_flows":20}
            previous={"lengths":[{"summary":summary}]}
            before=db.read_bytes()
            eda.analyze(root,db,catalog,previous,root,4266,100)
            report=json.loads((root/"results.json").read_text())
            self.assertEqual(report["metadata"]["raw_flows_reconciled"],240)
            self.assertEqual(len(report["candidates"]),8)
            self.assertEqual(db.read_bytes(),before)
            with np.load(root/"candidate_assignments.npz",allow_pickle=False) as manifest:
                self.assertEqual(len(manifest["start_ms"]),120)
                self.assertTrue(set(manifest["guarded_30m"].tolist())<={0,1,2,3})
            for r in report["isolation"]:
                self.assertEqual(r["retained_sequences"]+r["excluded_sequences"],120)
                if r["candidate"]!="random_100s":
                    self.assertEqual(r["shared_global_100s"],0)
                    self.assertEqual(r["interval_envelope_overlap_sequences"],0)
                if r["candidate"]=="guarded_30m":
                    self.assertEqual(r["cross_episodes_300s"],0)
            with self.assertRaises(ValueError):
                eda.load_catalog(catalog,{**summary,"flows":241})


if __name__=="__main__":
    unittest.main()
